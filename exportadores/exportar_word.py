"""Exportador de documentos a Word (.docx)"""

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
from datetime import datetime
import os


class ExportadorWord:
    """Clase para exportar el documento de glosa a Word"""
    
    def __init__(self, datos_proyecto: dict, contenidos: list, selecciones: dict):
        self.datos_proyecto = datos_proyecto
        self.contenidos = contenidos
        self.selecciones = selecciones
        self.documento = Document()
        self._configurar_documento()
    
    def _configurar_documento(self):
        """Configura los márgenes del documento"""
        secciones = self.documento.sections
        for seccion in secciones:
            seccion.top_margin = Cm(2.5)
            seccion.bottom_margin = Cm(2.5)
            seccion.left_margin = Cm(2.5)
            seccion.right_margin = Cm(2.5)
    
    def _agregar_encabezado(self):
        """Agrega el encabezado con título, fecha, código, proyecto y versión"""
        
        # Título principal centrado
        titulo = self.documento.add_paragraph()
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_titulo = titulo.add_run("TABLEROS DE DISTRIBUCIÓN ELÉCTRICA")
        run_titulo.bold = True
        run_titulo.font.size = Pt(16)
        run_titulo.font.name = 'Arial'
        
        # Subtítulo
        subtitulo = self.documento.add_paragraph()
        subtitulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = subtitulo.add_run("ESPECIFICACIONES TÉCNICAS")
        run_sub.font.size = Pt(12)
        run_sub.font.name = 'Arial'
        
        self.documento.add_paragraph()
        
        # Tabla con metadatos
        tabla_meta = self.documento.add_table(rows=1, cols=4)
        tabla_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        fecha_actual = datetime.now().strftime("%d/%m/%Y")
        nombre_proyecto = self.datos_proyecto.get('nombre', '')
        codigo_crm = self.datos_proyecto.get('codigo_crm', '')
        version = self.datos_proyecto.get('version', 'v1')
        
        if version:
            version_num = ''.join(filter(str.isdigit, version))
            version_formateada = f"V{version_num}" if version_num else "V1"
        else:
            version_formateada = "V1"
        
        celdas = tabla_meta.rows[0].cells
        celdas[0].text = f"Fecha: {fecha_actual}"
        celdas[1].text = f"Código: {codigo_crm}"
        celdas[2].text = f"Proyecto: {nombre_proyecto}"
        celdas[3].text = f"Versión: N°{version_formateada}"
        
        for celda in celdas:
            for paragraph in celda.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(9)
                    run.font.name = 'Arial'
        
        self.documento.add_paragraph()
    
    def _agregar_titulo_seccion_principal(self):
        """Agrega el título de la sección principal"""
        celda = self.selecciones.get("Celda", "")
        nivel = self.selecciones.get("Nivel de tensión", "")
        
        if celda in ["Seleccione las opciones anteriores", "Seleccione tipo, tensión y familia"]:
            celda = ""
        
        titulo_texto = "1. Celdas de media tensión"
        if celda and nivel:
            titulo_texto = f"1. Celdas de media tensión {celda} {nivel}"
        elif celda:
            titulo_texto = f"1. Celdas de media tensión {celda}"
        
        titulo = self.documento.add_paragraph()
        run = titulo.add_run(titulo_texto)
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = 'Arial'
        
        self.documento.add_paragraph()
    
    def _agregar_contenido(self):
        """Agrega todo el contenido del documento"""
        seccion_actual = ""
        
        for item in self.contenidos:
            seccion_num = item.get('seccion_numero', '')
            seccion_nom = item.get('seccion_nombre', '')
            
            if seccion_num and seccion_num != seccion_actual:
                seccion_actual = seccion_num
                
                titulo_seccion = self.documento.add_paragraph()
                run = titulo_seccion.add_run(f"{seccion_num} {seccion_nom}")
                run.bold = True
                run.font.size = Pt(12)
                run.font.name = 'Arial'
            
            tipo = item.get('tipo', 'texto')
            
            if tipo == 'texto':
                self._agregar_parrafo(item.get('texto', ''))
            
            elif tipo == 'lista':
                self._agregar_lista(item.get('texto', ''))
            
            elif tipo == 'imagen':
                imagen_data = item.get('imagen')
                if imagen_data:
                    self._agregar_imagen(imagen_data)
            
            elif tipo == 'tabla':
                tabla_data = item.get('tabla')
                if tabla_data:
                    self._agregar_tabla_word(item.get('titulo', ''), tabla_data)
    
    def _agregar_parrafo(self, texto: str):
        """Agrega un párrafo de texto"""
        if not texto:
            return
        
        parrafo = self.documento.add_paragraph()
        run = parrafo.add_run(texto)
        run.font.size = Pt(11)
        run.font.name = 'Arial'
        parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    def _agregar_lista(self, texto: str):
        """Agrega una lista con viñetas"""
        if not texto:
            return
        
        items = texto.split(';')
        for item in items:
            item = item.strip()
            if item:
                parrafo = self.documento.add_paragraph(style='List Bullet')
                run = parrafo.add_run(item)
                run.font.size = Pt(11)
                run.font.name = 'Arial'
    
    def _agregar_imagen(self, imagen_data: dict):
        """Agrega una imagen centrada con pie de imagen"""
        ruta_base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        ruta_imagen = os.path.join(ruta_base, "datos", imagen_data.get('ruta', ''))
        
        if os.path.exists(ruta_imagen):
            parrafo_img = self.documento.add_paragraph()
            parrafo_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = parrafo_img.add_run()
            
            ancho = min(imagen_data.get('ancho', 400) / 96, 4)
            run.add_picture(ruta_imagen, width=Inches(ancho))
            
            pie = imagen_data.get('pie', '')
            if pie:
                parrafo_pie = self.documento.add_paragraph()
                parrafo_pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_pie = parrafo_pie.add_run(pie)
                run_pie.bold = True
                run_pie.font.size = Pt(10)
                run_pie.font.name = 'Arial'
        
        self.documento.add_paragraph()
    
    def _agregar_tabla_word(self, titulo: str, tabla_data: dict):
        """Agrega una tabla al documento Word"""
        columnas = tabla_data.get('columnas', [])
        filas = tabla_data.get('filas', [])
        
        if not columnas or not filas:
            return
        
        # Título de la tabla
        if titulo:
            parrafo_titulo = self.documento.add_paragraph()
            parrafo_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = parrafo_titulo.add_run(titulo)
            run.bold = True
            run.font.size = Pt(10)
            run.font.name = 'Arial'
        
        # Crear tabla
        tabla = self.documento.add_table(rows=len(filas) + 1, cols=len(columnas))
        tabla.style = 'Table Grid'
        
        # Encabezados con fondo azul
        for col_idx, col_nombre in enumerate(columnas):
            celda = tabla.rows[0].cells[col_idx]
            celda.text = col_nombre
            
            # Fondo azul
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="0065bb"/>')
            celda._tc.get_or_add_tcPr().append(shading)
            
            for paragraph in celda.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(9)
                    run.font.name = 'Arial'
        
        # Datos
        for fila_idx, fila_datos in enumerate(filas):
            for col_idx, valor in enumerate(fila_datos):
                celda = tabla.rows[fila_idx + 1].cells[col_idx]
                celda.text = str(valor)
                
                for paragraph in celda.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        run.font.name = 'Arial'
        
        self.documento.add_paragraph()
    
    def exportar(self, ruta_archivo: str) -> bool:
        """Genera y guarda el documento Word"""
        try:
            self._agregar_encabezado()
            self._agregar_titulo_seccion_principal()
            self._agregar_contenido()
            
            self.documento.save(ruta_archivo)
            return True
        
        except Exception as e:
            print(f"Error al exportar a Word: {e}")
            return False


def exportar_glosa_a_word(datos_proyecto: dict, contenidos: list, selecciones: dict, ruta_archivo: str) -> bool:
    """Función auxiliar para exportar una glosa a Word"""
    exportador = ExportadorWord(datos_proyecto, contenidos, selecciones)
    return exportador.exportar(ruta_archivo)